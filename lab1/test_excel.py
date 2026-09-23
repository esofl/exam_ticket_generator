"""
Автоматизированные тесты для проверки функционала работы с Excel,
валидации ввода и логики приложения.
"""

from datetime import datetime
import os
import unittest
from unittest.mock import patch
from openpyxl import load_workbook

from console_input import prompt_non_empty
from excel_manager import HEADERS, append_record, init_journal_if_needed


TEST_FILE = "test_journal_tmp.xlsx"


class TestExamTicketGenerator(unittest.TestCase):
    def tearDown(self) -> None:
        if os.path.exists(TEST_FILE):
            try:
                os.remove(TEST_FILE)
            except OSError:
                pass

    def test_init_journal_creates_header(self) -> None:
        init_journal_if_needed(TEST_FILE)
        self.assertTrue(os.path.exists(TEST_FILE))

        wb = load_workbook(TEST_FILE)
        ws = wb.active
        self.assertEqual(ws.max_row, 1)

        row_values = [cell.value for cell in ws[1]]
        self.assertEqual(row_values, HEADERS)

    def test_append_record_preserves_previous_data(self) -> None:
        # Добавляем 1-го студента
        fixed_time1 = datetime(2026, 9, 23, 10, 0, 0)
        append_record("Петров", "Петр", 5, fixed_time1, filepath=TEST_FILE)

        # Добавляем 2-го студента
        fixed_time2 = datetime(2026, 9, 23, 10, 5, 0)
        append_record("Сидоров", "Алексей", 12, fixed_time2, filepath=TEST_FILE)

        wb = load_workbook(TEST_FILE)
        ws = wb.active
        self.assertEqual(ws.max_row, 3)

        row2 = [cell.value for cell in ws[2]]
        self.assertEqual(row2, ["Петров", "Петр", 5, "2026-09-23 10:00:00"])

        row3 = [cell.value for cell in ws[3]]
        self.assertEqual(row3, ["Сидоров", "Алексей", 12, "2026-09-23 10:05:00"])

    def test_ticket_range_validity(self) -> None:
        import random
        for _ in range(100):
            ticket = random.randint(1, 20)
            self.assertGreaterEqual(ticket, 1)
            self.assertLessEqual(ticket, 20)

    @patch("console_input.read_line_or_esc", side_effect=["", "   ", "Иванов"])
    def test_prompt_retries_on_empty(self, mock_read) -> None:
        result = prompt_non_empty("Last name: ")
        self.assertEqual(result, "Иванов")
        self.assertEqual(mock_read.call_count, 3)

    @patch("console_input.read_line_or_esc", side_effect=[None])
    def test_prompt_returns_none_on_esc(self, mock_read) -> None:
        result = prompt_non_empty("Last name: ")
        self.assertIsNone(result)


if __name__ == "__main__":
    unittest.main()
