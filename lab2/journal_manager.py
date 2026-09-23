"""
Модуль для работы с журналом результатов results.xlsx.
Обеспечивает создание журнала, поиск истории для соблюдения идемпотентности,
дозапись новых строк и обработку блокировок файла.
"""

from datetime import datetime
import logging
import os
from typing import Callable, List, Optional
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

from models import JournalRecord

logger = logging.getLogger(__name__)

HEADERS = ["Группа", "Фамилия", "Имя", "Номер билета", "Дата и время", "Повтор"]
DEFAULT_JOURNAL_FILE = "results.xlsx"


def init_journal_if_needed(filepath: str = DEFAULT_JOURNAL_FILE) -> None:
    """
    Создает файл results.xlsx со стандартными заголовками, если файл еще не существует.
    """
    if not os.path.exists(filepath):
        wb = Workbook()
        ws = wb.active
        ws.title = "Результаты"
        ws.append(HEADERS)

        # Стилизация шапки
        header_font = Font(bold=True, size=11)
        center_align = Alignment(horizontal="center", vertical="center")
        for col_idx in range(1, len(HEADERS) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.alignment = center_align

        # Установка оптимальной ширины колонок
        column_widths = {
            "A": 16,  # Группа
            "B": 20,  # Фамилия
            "C": 20,  # Имя
            "D": 16,  # Номер билета
            "E": 22,  # Дата и время
            "F": 12,  # Повтор
        }
        for col_letter, width in column_widths.items():
            ws.column_dimensions[col_letter].width = width

        wb.save(filepath)
        logger.info(f"Создан новый файл журнала: {filepath}")


def get_first_ticket_for_student(
    filepath: str,
    group: str,
    last_name: str,
    first_name: str
) -> Optional[int]:
    """
    Ищет номер билета из самой первой записи по студенту в results.xlsx.
    Сравнение регистронезависимое с обрезкой пробелов.

    :return: Номер билета (int) или None, если студент еще не записан в журнал.
    """
    if not os.path.exists(filepath):
        return None

    norm_group = group.strip().lower()
    norm_last = last_name.strip().lower()
    norm_first = first_name.strip().lower()

    try:
        wb = load_workbook(filepath, read_only=True, data_only=True)
    except Exception as e:
        logger.error(f"Ошибка чтения файла журнала '{filepath}': {e}")
        return None

    try:
        ws = wb.active
        # Сканируем строки сверху вниз (начиная со 2-й строки)
        for row in ws.iter_rows(min_row=2, max_col=6, values_only=True):
            if not row or len(row) < 4:
                continue
            r_group, r_last, r_first, r_ticket = row[0], row[1], row[2], row[3]
            if r_group is None or r_last is None or r_first is None or r_ticket is None:
                continue

            if (
                str(r_group).strip().lower() == norm_group
                and str(r_last).strip().lower() == norm_last
                and str(r_first).strip().lower() == norm_first
            ):
                try:
                    return int(r_ticket)
                except ValueError:
                    logger.warning(f"Не удалось преобразовать номер билета '{r_ticket}' к числу.")
                    continue
        return None
    finally:
        wb.close()


def append_result_record(
    group: str,
    last_name: str,
    first_name: str,
    ticket_number: int,
    is_repeat: bool,
    timestamp: Optional[datetime] = None,
    filepath: str = DEFAULT_JOURNAL_FILE,
    on_lock_retry: Optional[Callable[[], bool]] = None,
) -> None:
    """
    Дописывает новую запись в конец results.xlsx и сохраняет файл немедленно.

    :param group: Имя группы
    :param last_name: Фамилия студента
    :param first_name: Имя студента
    :param ticket_number: Номер билета
    :param is_repeat: Флаг повтора (True -> 'да', False -> 'нет')
    :param timestamp: Дата и время генерации (по умолчанию текущее время)
    :param filepath: Путь к файлу results.xlsx
    :param on_lock_retry: Функция обратного вызова при PermissionError. Должна возвращать
                          True для повтора попытки, False для отмены.
    """
    init_journal_if_needed(filepath)

    if timestamp is None:
        timestamp = datetime.now()
    formatted_time = timestamp.strftime("%Y-%m-%d %H:%M:%S")
    repeat_str = "да" if is_repeat else "нет"

    row_data = [group.strip(), last_name.strip(), first_name.strip(), ticket_number, formatted_time, repeat_str]

    while True:
        try:
            wb = load_workbook(filepath)
            ws = wb.active
            ws.append(row_data)

            # Выравнивание по центру для номера билета, даты и флага повтора
            last_row = ws.max_row
            center_align = Alignment(horizontal="center", vertical="center")
            ws.cell(row=last_row, column=4).alignment = center_align
            ws.cell(row=last_row, column=5).alignment = center_align
            ws.cell(row=last_row, column=6).alignment = center_align

            wb.save(filepath)
            logger.info(
                f"Запись сохранена в '{filepath}': {group} | {last_name} {first_name} | "
                f"Билет № {ticket_number} | Повтор: {repeat_str}"
            )
            return
        except PermissionError as pe:
            logger.warning(f"Файл '{filepath}' заблокирован другим процессом: {pe}")
            if on_lock_retry is not None:
                should_retry = on_lock_retry()
                if should_retry:
                    continue
                else:
                    raise PermissionError(f"Запись отменена пользователем. Файл '{filepath}' заблокирован.") from pe
            else:
                # В неинтерактивном окружении пробрасываем исключение
                raise


def get_all_records(filepath: str = DEFAULT_JOURNAL_FILE) -> List[JournalRecord]:
    """Считывает все записи из журнала results.xlsx (используется для тестов)."""
    if not os.path.exists(filepath):
        return []

    records: List[JournalRecord] = []
    wb = load_workbook(filepath, read_only=True, data_only=True)
    try:
        ws = wb.active
        for row in ws.iter_rows(min_row=2, max_col=6, values_only=True):
            if not row or len(row) < 6:
                continue
            r_group, r_last, r_first, r_ticket, r_time, r_rep = row[:6]
            if r_group is None:
                continue
            records.append(
                JournalRecord(
                    group=str(r_group),
                    last_name=str(r_last),
                    first_name=str(r_first),
                    ticket_number=int(r_ticket),
                    timestamp_str=str(r_time),
                    is_repeat=(str(r_rep).strip().lower() == "да")
                )
            )
        return records
    finally:
        wb.close()
