"""
Модуль чтения групп и списков студентов из файла students.xlsx.
"""

import logging
import os
from typing import Dict, List
from openpyxl import load_workbook

from models import Student

logger = logging.getLogger(__name__)


def get_group_names(filepath: str) -> List[str]:
    """
    Возвращает список групп (имен листов) из students.xlsx.

    :param filepath: Путь к файлу students.xlsx
    :return: Список имен групп
    :raises FileNotFoundError: Если файл не найден
    """
    if not os.path.exists(filepath):
        raise FileNotFoundError(f"Файл со студентами '{filepath}' не найден.")

    wb = load_workbook(filepath, read_only=True)
    try:
        groups = [str(name).strip() for name in wb.sheetnames if str(name).strip()]
        return groups
    finally:
        wb.close()


def get_students_for_group(filepath: str, group_name: str) -> List[Student]:
    """
    Считывает список студентов для указанной группы из соответствующего листа.

    Формат листа:
    Строка 1: Шапка (Колонка A — Фамилия, Колонка B — Имя)
    Строка 2+: Данные студентов.

    Если группа пуста (нет строк данных), возвращается пустой список без исключений.

    :param filepath: Путь к файлу students.xlsx
    :param group_name: Имя группы (название листа)
    :return: Список объектов Student
    """
    if not os.path.exists(filepath):
        raise FileNotFoundError(f"Файл со студентами '{filepath}' не найден.")

    wb = load_workbook(filepath, read_only=True, data_only=True)
    students: List[Student] = []

    try:
        if group_name not in wb.sheetnames:
            logger.warning(f"Лист группы '{group_name}' не найден в файле '{filepath}'.")
            return []

        ws = wb[group_name]
        # Итерируемся начиная со 2-й строки
        for row in ws.iter_rows(min_row=2, max_col=2, values_only=True):
            if not row or len(row) < 2:
                continue
            last_name_val, first_name_val = row[0], row[1]
            if last_name_val is None and first_name_val is None:
                continue

            last_name = str(last_name_val).strip() if last_name_val is not None else ""
            first_name = str(first_name_val).strip() if first_name_val is not None else ""

            # Игнорируем строки, где оба поля пустые
            if not last_name and not first_name:
                continue

            students.append(
                Student(
                    group=group_name,
                    last_name=last_name,
                    first_name=first_name
                )
            )

        logger.info(f"Загружено {len(students)} студентов для группы '{group_name}'.")
        return students
    finally:
        wb.close()
