"""
Автоматизированные тесты для Лабораторной работы №2.
Проверяют:
1. Парсинг билетов из tickets.docx (включая пропуск билетов с < 3 вопросов).
2. Считывание групп и студентов из students.xlsx (включая пустые группы).
3. Идемпотентность назначения билетов и сохранение истории в results.xlsx.
4. Немедленную дозапись в results.xlsx без изменения предыдущих строк.
5. Обработку блокировки файла и отсутствующих источников.
"""

from datetime import datetime
import os
import unittest
from unittest.mock import MagicMock, patch

from create_sample_data import create_sample_students_xlsx, create_sample_tickets_docx
from docx_parser import parse_tickets_from_docx
from excel_reader import get_group_names, get_students_for_group
from journal_manager import (
    DEFAULT_JOURNAL_FILE,
    HEADERS,
    append_result_record,
    get_all_records,
    get_first_ticket_for_student,
    init_journal_if_needed,
)
from models import Student
from ticket_service import TicketService

TEST_STUDENTS_FILE = "test_students_tmp.xlsx"
TEST_TICKETS_FILE = "test_tickets_tmp.docx"
TEST_RESULTS_FILE = "test_results_tmp.xlsx"


class TestLab2TicketGenerator(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        create_sample_students_xlsx(TEST_STUDENTS_FILE)
        create_sample_tickets_docx(TEST_TICKETS_FILE)

    @classmethod
    def tearDownClass(cls):
        for f in (TEST_STUDENTS_FILE, TEST_TICKETS_FILE):
            if os.path.exists(f):
                try:
                    os.remove(f)
                except OSError:
                    pass

    def tearDown(self):
        if os.path.exists(TEST_RESULTS_FILE):
            try:
                os.remove(TEST_RESULTS_FILE)
            except OSError:
                pass

    # --- 1. ТЕСТЫ ПАРСИНГА WORD ---

    def test_parse_docx_valid_tickets_and_skip_invalid(self):
        """Парсинг должен успешно загрузить билеты 1-5 и пропустить билет 99 (<3 вопросов)."""
        tickets = parse_tickets_from_docx(TEST_TICKETS_FILE)

        # Билеты 1..5 должны быть загружены
        for num in range(1, 6):
            self.assertIn(num, tickets)
            self.assertEqual(len(tickets[num].questions), 3)

        # Билет 99 имел только 1 вопрос и должен быть пропущен
        self.assertNotIn(99, tickets)

    def test_parse_docx_missing_file_raises_not_found(self):
        with self.assertRaises(FileNotFoundError):
            parse_tickets_from_docx("non_existent_file.docx")

    # --- 2. ТЕСТЫ ЧТЕНИЯ EXCEL (СТУДЕНТЫ И ГРУППЫ) ---

    def test_excel_reader_groups(self):
        groups = get_group_names(TEST_STUDENTS_FILE)
        self.assertIn("ИТ-21", groups)
        self.assertIn("ПИ-22", groups)
        self.assertIn("Пустая-Группа", groups)

    def test_excel_reader_students(self):
        students_it21 = get_students_for_group(TEST_STUDENTS_FILE, "ИТ-21")
        self.assertEqual(len(students_it21), 5)
        self.assertEqual(students_it21[0].last_name, "Иванов")
        self.assertEqual(students_it21[0].first_name, "Иван")
        self.assertEqual(students_it21[0].full_name, "Иванов Иван")

    def test_excel_reader_empty_group_returns_empty_list_without_exception(self):
        """Пустая группа не должна вызывать исключений (требование ТЗ)."""
        students_empty = get_students_for_group(TEST_STUDENTS_FILE, "Пустая-Группа")
        self.assertEqual(students_empty, [])

    # --- 3. ТЕСТЫ ЖУРНАЛА И ИДЕМПОТЕНТНОСТИ ---

    def test_init_journal_creates_correct_headers(self):
        init_journal_if_needed(TEST_RESULTS_FILE)
        self.assertTrue(os.path.exists(TEST_RESULTS_FILE))

        from openpyxl import load_workbook
        wb = load_workbook(TEST_RESULTS_FILE)
        ws = wb.active
        self.assertEqual(ws.max_row, 1)
        row_values = [cell.value for cell in ws[1]]
        self.assertEqual(row_values, HEADERS)

    def test_idempotency_ticket_assignment(self):
        """
        Проверка главного правила назначения билетов:
        - 1-й выбор студента: случайный билет, Повтор = 'нет'.
        - 2-й выбор того же студента: ТОТ ЖЕ номер билета, Повтор = 'да'.
        - Новая строка добавляется в конец журнала, старые не затираются.
        """
        service = TicketService(
            students_filepath=TEST_STUDENTS_FILE,
            tickets_filepath=TEST_TICKETS_FILE,
            results_filepath=TEST_RESULTS_FILE,
        )

        student1 = Student(group="ИТ-21", last_name="Иванов", first_name="Иван")

        # 1. Первая генерация
        ticket1, is_repeat1 = service.generate_or_get_ticket(student1)
        self.assertFalse(is_repeat1)
        self.assertIn(ticket1.number, range(1, 6))

        # 2. Повторная генерация для того же студента
        ticket2, is_repeat2 = service.generate_or_get_ticket(student1)
        self.assertTrue(is_repeat2)
        self.assertEqual(ticket1.number, ticket2.number, "Номер билета при повторе должен совпадать!")

        # 3. Генерация для другого студента
        student2 = Student(group="ИТ-21", last_name="Петров", first_name="Петр")
        ticket3, is_repeat3 = service.generate_or_get_ticket(student2)
        self.assertFalse(is_repeat3)

        # 4. Проверяем историю в файле results.xlsx
        records = get_all_records(TEST_RESULTS_FILE)
        self.assertEqual(len(records), 3, "В журнале должно быть ровно 3 записи")

        # Запись 1: Иванов, повтор = False
        self.assertEqual(records[0].last_name, "Иванов")
        self.assertEqual(records[0].ticket_number, ticket1.number)
        self.assertFalse(records[0].is_repeat)

        # Запись 2: Иванов, повтор = True, номер тот же
        self.assertEqual(records[1].last_name, "Иванов")
        self.assertEqual(records[1].ticket_number, ticket1.number)
        self.assertTrue(records[1].is_repeat)

        # Запись 3: Петров, повтор = False
        self.assertEqual(records[2].last_name, "Петров")
        self.assertEqual(records[2].ticket_number, ticket3.number)
        self.assertFalse(records[2].is_repeat)

    def test_file_lock_retry_callback(self):
        """Проверка реакции на блокировку файла PermissionError с вызовом retry."""
        retry_called = 0

        def on_retry():
            nonlocal retry_called
            retry_called += 1
            return False  # Отмена после первой попытки

        with patch("journal_manager.load_workbook", side_effect=PermissionError("File locked")):
            with self.assertRaises(PermissionError):
                append_result_record(
                    group="ИТ-21",
                    last_name="Тестов",
                    first_name="Тест",
                    ticket_number=1,
                    is_repeat=False,
                    filepath=TEST_RESULTS_FILE,
                    on_lock_retry=on_retry
                )

        self.assertEqual(retry_called, 1, "Callback повтора должен быть вызван при блокировке")


if __name__ == "__main__":
    unittest.main()
