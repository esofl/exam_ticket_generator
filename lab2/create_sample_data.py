"""
Скрипт создания демонстрационных файлов students.xlsx и tickets.docx
для проверки функционала в точном соответствии с ТЗ Лабораторной работы №2.
"""

import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
import docx


def create_sample_students_xlsx(filepath: str = "students.xlsx") -> None:
    wb = Workbook()
    
    # Лист 1: ИТ-21
    ws1 = wb.active
    ws1.title = "ИТ-21"
    ws1.append(["Фамилия", "Имя"])
    students_it21 = [
        ("Иванов", "Иван"),
        ("Петров", "Петр"),
        ("Сидоров", "Алексей"),
        ("Смирнова", "Анна"),
        ("Кузнецов", "Дмитрий"),
    ]
    for row in students_it21:
        ws1.append(list(row))

    # Лист 2: ПИ-22
    ws2 = wb.create_sheet(title="ПИ-22")
    ws2.append(["Фамилия", "Имя"])
    students_pi22 = [
        ("Васильев", "Михаил"),
        ("Михайлова", "Елена"),
        ("Федоров", "Сергей"),
        ("Морозова", "Ольга"),
    ]
    for row in students_pi22:
        ws2.append(list(row))

    # Лист 3: Пустая-Группа (для проверки устойчивости по ТЗ)
    ws3 = wb.create_sheet(title="Пустая-Группа")
    ws3.append(["Фамилия", "Имя"])

    # Стилизация шапок
    header_font = Font(bold=True)
    center_align = Alignment(horizontal="center")
    for ws in [ws1, ws2, ws3]:
        for col_idx in (1, 2):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.alignment = center_align
        ws.column_dimensions["A"].width = 20
        ws.column_dimensions["B"].width = 20

    wb.save(filepath)
    print(f"[OK] Создан демонстрационный файл: {filepath}")


def create_sample_tickets_docx(filepath: str = "tickets.docx") -> None:
    doc = docx.Document()

    tickets_data = [
        (
            1,
            [
                "Что такое полиморфизм в объектно-ориентированном программировании?",
                "Объясните назначение и алгоритм работы сборщика мусора (Garbage Collector).",
                "В чем фундаментальное отличие процесса от потока операционной системы?",
            ]
        ),
        (
            2,
            [
                "Принципы SOLID: подробно опишите принцип открытости/закрытости (OCP).",
                "Что такое deadlock (взаимная блокировка) и каковы условия его возникновения?",
                "Как устроена хеш-таблица и как разрешаются коллизии?",
            ]
        ),
        (
            3,
            [
                "Различия между протоколами TCP и UDP на транспортном уровне модели OSI.",
                "Паттерн проектирования Singleton: достоинства, недостатки и многопоточная реализация.",
                "Что такое индексы в реляционных базах данных и как они влияют на производительность?",
            ]
        ),
        (
            4,
            [
                "Архитектурный шаблон MVC: назначение компонентов Model, View, Controller.",
                "Что такое транзакция в СУБД? Опишите свойства ACID.",
                "Стек вызовов (call stack) и куча (heap): распределение памяти в программах.",
            ]
        ),
        (
            5,
            [
                "Принцип инверсии зависимостей (Dependency Inversion Principle) и Dependency Injection.",
                "Что такое состояние гонки (race condition) и способы его предотвращения.",
                "Разница между симметричным и асимметричным шифрованием.",
            ]
        ),
    ]

    for t_num, questions in tickets_data:
        doc.add_paragraph(f"Билет {t_num}")
        for q_idx, q_text in enumerate(questions, start=1):
            doc.add_paragraph(f"{q_idx}. {q_text}")
        doc.add_paragraph("")  # Пустая строка-разделитель

    # Добавляем 1 намеренно невалидный билет (<3 вопросов) для проверки отказоустойчивости по ТЗ
    doc.add_paragraph("Билет 99")
    doc.add_paragraph("1. Единственный вопрос неполного билета (должен быть пропущен парсером).")
    doc.add_paragraph("")

    doc.save(filepath)
    print(f"[OK] Создан демонстрационный файл: {filepath}")


if __name__ == "__main__":
    create_sample_students_xlsx("students.xlsx")
    create_sample_tickets_docx("tickets.docx")
    print("Генерация демо-данных завершена успешно.")
