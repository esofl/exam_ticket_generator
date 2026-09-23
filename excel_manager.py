"""
Модуль для работы с файлом Excel (journal.xlsx).
Обеспечивает создание файла с шапкой, дозапись строк и обработку блокировок файла.
"""

from datetime import datetime
import os
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font


HEADERS = ["Last name", "First name", "Номер билета", "Дата и время"]
DEFAULT_FILENAME = "journal.xlsx"


def init_journal_if_needed(filepath: str = DEFAULT_FILENAME) -> None:
    """
    Создает файл журнала с шапкой, если он еще не существует.
    Если файл уже есть — ничего не перезаписывает.
    """
    if not os.path.exists(filepath):
        wb = Workbook()
        ws = wb.active
        ws.title = "Журнал"
        ws.append(HEADERS)

        # Оформление шапки
        header_font = Font(bold=True)
        center_align = Alignment(horizontal="center", vertical="center")
        for col_idx in range(1, len(HEADERS) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.alignment = center_align

        # Установка удобной ширины колонок
        column_widths = {"A": 18, "B": 18, "C": 16, "D": 22}
        for col_letter, width in column_widths.items():
            ws.column_dimensions[col_letter].width = width

        _save_with_retry(wb, filepath)


def append_record(
    last_name: str,
    first_name: str,
    ticket_number: int,
    timestamp: datetime | None = None,
    filepath: str = DEFAULT_FILENAME,
) -> None:
    """
    Дописывает запись о студенте в конец файла journal.xlsx.
    Сохраняет файл сразу же после добавления записи.
    """
    init_journal_if_needed(filepath)

    if timestamp is None:
        timestamp = datetime.now()
    formatted_time = timestamp.strftime("%Y-%m-%d %H:%M:%S")

    while True:
        try:
            wb = load_workbook(filepath)
            ws = wb.active

            row_data = [last_name, first_name, ticket_number, formatted_time]
            ws.append(row_data)

            # Выравнивание номера билета и даты по центру
            last_row = ws.max_row
            ws.cell(row=last_row, column=3).alignment = Alignment(horizontal="center")
            ws.cell(row=last_row, column=4).alignment = Alignment(horizontal="center")

            _save_with_retry(wb, filepath)
            break
        except PermissionError:
            print(f"\n[ВНИМАНИЕ] Файл '{filepath}' открыт в другой программе (например, Microsoft Excel).")
            print("Пожалуйста, закройте файл в Excel и нажмите Enter для повторной попытки сохранения...")
            input()
        except Exception as e:
            print(f"\n[ОШИБКА] Не удалось записать в '{filepath}': {e}")
            print("Нажмите Enter для повторной попытки...")
            input()


def _save_with_retry(wb: Workbook, filepath: str) -> None:
    """
    Сохраняет книгу Excel с обработкой блокировки файла (PermissionError).
    Предлагает пользователю закрыть файл и нажать Enter, предотвращая падение программы.
    """
    while True:
        try:
            wb.save(filepath)
            return
        except PermissionError:
            print(f"\n[ВНИМАНИЕ] Файл '{filepath}' открыт в другой программе (например, Microsoft Excel).")
            print("Пожалуйста, закройте файл в Excel и нажмите Enter для повторной попытки сохранения...")
            input()
